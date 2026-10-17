"""Sampled interventional Shapley values for SSI graph/recurrent model inputs.

Each feature is one named channel. Node, element and earthquake-window channels
are explained as groups across all nodes, edges or time windows. The baseline is
the training-fold mean (zero after standardization); graph connectivity is held
fixed because it is topology, not a scalar input feature.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import h5py
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook


def _json_attr(h5, key):
    raw = h5.attrs.get(key, "[]")
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8", errors="replace")
    try:
        return json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return []


def feature_groups(database: str, first_case: str):
    """Return (input dictionary key, names) groups aligned with the model tensors."""
    metadata = {}
    sidecar = Path(database).with_suffix(Path(database).suffix + ".metadata.json")
    if sidecar.is_file():
        try:
            metadata = json.loads(sidecar.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            metadata = {}
    ordered = metadata.get("ordered_feature_names", {})
    with h5py.File(database, "r") as h:
        c = h[first_case]
        eq_global = _json_attr(h, "global_eq_feature_names_json")
        eq_window = _json_attr(h, "window_feature_names_json")
        static_names = ordered.get("global_bridge_soil_input_features", [])
        node_names = ordered.get("node_features", [])
        edge_names = ordered.get("edge_features", [])
        window_names = ordered.get("earthquake_window_sequence_features", eq_window)
        dims = {
            "global_features": int(c["static"].shape[0] + c["eq_global"].shape[0]),
            "node_features": int(c["node_features"].shape[1]),
            "edge_features": int(c["edge_features"].shape[1]),
            "ground_motion": int(c["eq_sequence"].shape[1]),
        }
    named = {
        "global_features": list(static_names) + [f"earthquake_global.{x}" for x in eq_global],
        "node_features": list(node_names),
        "edge_features": list(edge_names),
        "ground_motion": list(window_names),
    }
    output = []
    for key, width in dims.items():
        names = named[key]
        if len(names) != width:
            names = [f"{key}_feature_{i:02d}" for i in range(width)]
        output.extend((key, i, str(name)) for i, name in enumerate(names))
    return output


def sampled_shapley(model, context, groups, output_scale, permutations, rng):
    """Estimate feature Shapley values by averaging random feature orderings."""
    keys = ("global_features", "node_features", "edge_features", "ground_motion")
    original = {key: context[key] for key in keys}
    fixed = {key: value for key, value in context.items() if key not in keys}
    n_features = len(groups)
    values = np.zeros((n_features, 2), dtype=np.float64)
    scales = np.asarray(output_scale, dtype=np.float64).reshape(1, 2)

    def predict(active):
        state = dict(fixed)
        state.update(active)
        with __import__("torch").no_grad():
            return model(state).detach().cpu().numpy().astype(np.float64) * scales

    for _ in range(permutations):
        active = {key: value.new_zeros(value.shape) for key, value in original.items()}
        previous = predict(active)
        for index in rng.permutation(n_features):
            key, column, _ = groups[int(index)]
            active[key][..., column].copy_(original[key][..., column])
            current = predict(active)
            values[int(index)] += (current - previous).reshape(2)
            previous = current
    return (values / float(permutations)).astype(np.float32)


def feature_value_summaries(case, groups):
    """Summarize each raw (unstandardized) channel used by one case.

    Global channels are scalar values. Node, edge, and earthquake-window
    channels are groups; their mean/std/min/max over positions are exported so
    users can plot a representative x value against the grouped Shapley value.
    """
    arrays = {
        "global_features": np.asarray(case[0]),
        "node_features": np.asarray(case[1]),
        "edge_features": np.asarray(case[3]),
        "ground_motion": np.asarray(case[4]),
    }
    summaries = []
    for key, column, _ in groups:
        values = np.asarray(arrays[key][..., column], dtype=np.float64).reshape(-1)
        values = values[np.isfinite(values)]
        if values.size == 0:
            summaries.append((float("nan"),) * 4)
        else:
            summaries.append((float(values.mean()), float(values.std()),
                              float(values.min()), float(values.max())))
    return summaries


def _summaries(model_samples, names):
    rows = []
    model_names = sorted(model_samples)
    for model in model_names:
        records = model_samples[model]
        fold_ids = sorted({int(record[0]) for record in records})
        for fold in fold_ids:
            subset = [r[2] for r in records if int(r[0]) == fold]
            arr = np.stack(subset)
            for output_index, output in enumerate(("Ux", "Uy")):
                for feature_index, feature in enumerate(names):
                    v = arr[:, feature_index, output_index]
                    rows.append({"model": model, "fold": fold, "output": output,
                        "feature": feature, "mean_abs_shapley": float(np.mean(np.abs(v))),
                        "mean_signed_shapley": float(np.mean(v)), "sample_count": len(v)})
        arr = np.stack([r[2] for r in records])
        for output_index, output in enumerate(("Ux", "Uy")):
            for feature_index, feature in enumerate(names):
                v = arr[:, feature_index, output_index]
                rows.append({"model": model, "fold": "ALL", "output": output,
                    "feature": feature, "mean_abs_shapley": float(np.mean(np.abs(v))),
                    "mean_signed_shapley": float(np.mean(v)), "sample_count": len(v)})
    return rows


def write_outputs(output_dir: Path, model_samples, groups, permutations):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    names = [group[2] for group in groups]
    rows = _summaries(model_samples, names)
    csv_path = output_dir / "shapley_feature_importance.csv"
    columns = ["model", "fold", "output", "feature", "mean_abs_shapley",
               "mean_signed_shapley", "sample_count"]
    with csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)

    sample_columns = ["fold", "simulation", "model", "output", "feature",
                      "input_feature_mean", "input_feature_std", "input_feature_min",
                      "input_feature_max", "shapley_value", "permutations"]
    sample_rows = []
    for model, records in model_samples.items():
        for fold, simulation, phi, feature_values in records:
            for feature_index, (_, _, feature) in enumerate(groups):
                mean, std, minimum, maximum = feature_values[feature_index]
                for output_index, output in enumerate(("Ux", "Uy")):
                    sample_rows.append({"fold": fold, "simulation": simulation,
                        "model": model, "output": output, "feature": feature,
                        "input_feature_mean": mean, "input_feature_std": std,
                        "input_feature_min": minimum, "input_feature_max": maximum,
                        "shapley_value": float(phi[feature_index, output_index]),
                        "permutations": permutations})
    sample_csv_path = output_dir / "shapley_sample_values.csv"
    with sample_csv_path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=sample_columns)
        writer.writeheader()
        writer.writerows(sample_rows)

    book = Workbook()
    all_ws = book.active
    all_ws.title = "All_Features"
    all_ws.append(columns + ["permutations_per_case", "interpretation"])
    note = "Mean absolute sampled Shapley contribution; baseline is training-fold mean; graph topology fixed."
    for row in rows:
        all_ws.append([row[c] for c in columns] + [permutations, note])
    all_ws.freeze_panes = "A2"
    all_ws.auto_filter.ref = all_ws.dimensions
    rank_ws = book.create_sheet("Top_Features")
    rank_ws.append(["model", "fold", "output", "rank", "feature", "mean_abs_shapley", "mean_signed_shapley", "sample_count"])
    for model in sorted(model_samples):
        for output in ("Ux", "Uy"):
            chosen = [r for r in rows if r["model"] == model and r["fold"] == "ALL" and r["output"] == output]
            chosen.sort(key=lambda row: row["mean_abs_shapley"], reverse=True)
            for rank, row in enumerate(chosen, 1):
                rank_ws.append([model, "ALL", output, rank, row["feature"], row["mean_abs_shapley"], row["mean_signed_shapley"], row["sample_count"]])
    rank_ws.freeze_panes = "A2"
    rank_ws.auto_filter.ref = rank_ws.dimensions
    sample_ws = book.create_sheet("Sampled_SHAP_Values")
    sample_ws.append(sample_columns)
    for row in sample_rows:
        sample_ws.append([row[c] for c in sample_columns])
    sample_ws.freeze_panes = "A2"
    if sample_ws.max_row > 1:
        sample_ws.auto_filter.ref = sample_ws.dimensions
    xlsx_path = output_dir / "shapley_feature_importance.xlsx"
    book.save(xlsx_path)

    model_names = sorted(model_samples)
    palette = plt.get_cmap("tab10")
    # Each figure contains at most six feature names. Further features are
    # written to the following numbered figure, separately for Ux and Uy.
    for output in ("Ux", "Uy"):
        all_rows = [r for r in rows if r["fold"] == "ALL" and r["output"] == output]
        rank_by_feature = {}
        for row in all_rows:
            rank_by_feature[row["feature"]] = max(
                rank_by_feature.get(row["feature"], 0.0), row["mean_abs_shapley"])
        ranked_features = [name for name, _ in sorted(rank_by_feature.items(), key=lambda item: item[1], reverse=True)]
        for page, start in enumerate(range(0, len(ranked_features), 6), 1):
            page_features = ranked_features[start:start + 6]
            fig, ax = plt.subplots(figsize=(12, max(4, .65 * len(page_features) + 1.5)))
            y = np.arange(len(page_features), dtype=float)
            height = .8 / max(1, len(model_names))
            for mi, model in enumerate(model_names):
                vals = []
                for feature in page_features:
                    found = next((r["mean_abs_shapley"] for r in all_rows
                                  if r["model"] == model and r["feature"] == feature), 0.0)
                    vals.append(found)
                offset = (mi - (len(model_names) - 1) / 2) * height
                ax.barh(y + offset, vals, height=height, color=palette(mi), label=model)
            ax.set_yticks(y, page_features)
            ax.invert_yaxis()
            ax.set_xlabel(f"Mean |Shapley value| for {output} peak displacement (m)")
            ax.set_title(f"Top feature influence on {output}: features {start + 1}–{start + len(page_features)}")
            ax.grid(axis="x", alpha=.25)
            ax.legend(fontsize=8, ncol=2)
            fig.tight_layout()
            fig.savefig(output_dir / f"shapley_{output}_features_{page:02d}.png", dpi=160)
            plt.close(fig)
    print(f"Shapley outputs for {len(names)} named input channels: {csv_path}, {sample_csv_path}, {xlsx_path}, and per-output plots in {output_dir}", flush=True)
